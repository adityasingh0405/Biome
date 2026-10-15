import os
import sys
import docx
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from styles import (
    configure_document_styles, set_page_margins, set_section_pagination,
    add_page_number
)
from section_prelims import (
    add_cover_page, add_student_declaration, add_offer_letter,
    add_completion_certificate, add_employer_feedback,
    add_acknowledgement, add_table_of_contents_view,
    add_list_of_figures, add_list_of_tables
)
from section_chapter1 import build_chapter_1
from section_chapter2 import build_chapter_2
from section_chapter3 import build_chapter_3
from section_chapter4 import build_chapter_4
from section_post import build_references_and_appendices

OUTPUT_DOCX = "Summer_Training _Report_Format.docx"

def attach_footer_page_number(footer, alignment=WD_PARAGRAPH_ALIGNMENT.CENTER):
    """Configures footer with bottom-center page number field."""
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.text = ""
    p.alignment = alignment
    run = p.add_run()
    run.font.name = 'Times New Roman'
    run.font.size = docx.shared.Pt(11)
    add_page_number(run)

def main():
    print("Initializing document generation for Biome Enterprise RAG...")
    doc = docx.Document()
    configure_document_styles(doc)
    
    # -------------------------------------------------------------
    # SECTION 1: COVER PAGE
    # -------------------------------------------------------------
    sec_cover = doc.sections[0]
    set_page_margins(sec_cover)
    sec_cover.different_first_page_header_footer = True
    # Unlinked, empty footer
    sec_cover.first_page_footer.is_linked_to_previous = False
    sec_cover.first_page_footer.paragraphs[0].text = ""
    
    print("Generating Cover Page...")
    add_cover_page(doc)
    
    # -------------------------------------------------------------
    # SECTION 2: PRELIMINARY PAGES (Roman Numerals: i, ii, iii...)
    # -------------------------------------------------------------
    sec_prelim = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
    set_page_margins(sec_prelim)
    set_section_pagination(sec_prelim, fmt="lowerRoman", start_num=1, restart=True)
    attach_footer_page_number(sec_prelim.footer)
    
    print("Generating Candidate Declaration...")
    add_student_declaration(doc)
    doc.add_page_break()
    
    print("Generating Offer Letter...")
    add_offer_letter(doc)
    doc.add_page_break()
    
    print("Generating Company Completion Certificate...")
    add_completion_certificate(doc)
    doc.add_page_break()
    
    print("Generating Employer Feedback...")
    add_employer_feedback(doc)
    doc.add_page_break()
    
    print("Generating Acknowledgement...")
    add_acknowledgement(doc)
    doc.add_page_break()
    
    print("Generating Table of Contents...")
    add_table_of_contents_view(doc)
    doc.add_page_break()
    
    print("Generating List of Figures...")
    add_list_of_figures(doc)
    doc.add_page_break()
    
    print("Generating List of Tables...")
    add_list_of_tables(doc)
    
    # -------------------------------------------------------------
    # SECTION 3: CHAPTER 1 (Arabic Numerals: 1, 2, 3... restarts at 1)
    # -------------------------------------------------------------
    sec_ch1 = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
    set_page_margins(sec_ch1)
    sec_ch1.different_first_page_header_footer = True
    set_section_pagination(sec_ch1, fmt="decimal", start_num=1, restart=True)
    # First page footer empty (accounted for but hidden)
    sec_ch1.first_page_footer.is_linked_to_previous = False
    sec_ch1.first_page_footer.paragraphs[0].text = ""
    # Subsequent pages print Arabic numeral
    attach_footer_page_number(sec_ch1.footer)
    
    print("Generating Chapter 1: Introduction...")
    build_chapter_1(doc)
    
    # -------------------------------------------------------------
    # SECTION 4: CHAPTER 2 (Arabic Numerals: Continues from Ch 1)
    # -------------------------------------------------------------
    sec_ch2 = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
    set_page_margins(sec_ch2)
    sec_ch2.different_first_page_header_footer = True
    set_section_pagination(sec_ch2, fmt="decimal", restart=False)
    # First page footer empty
    sec_ch2.first_page_footer.is_linked_to_previous = False
    sec_ch2.first_page_footer.paragraphs[0].text = ""
    # Subsequent pages continue Arabic numerals
    attach_footer_page_number(sec_ch2.footer)
    
    print("Generating Chapter 2: Project Architecture & System Design...")
    build_chapter_2(doc)
    
    # -------------------------------------------------------------
    # SECTION 5: CHAPTER 3 (Arabic Numerals: Continues from Ch 2)
    # -------------------------------------------------------------
    sec_ch3 = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
    set_page_margins(sec_ch3)
    sec_ch3.different_first_page_header_footer = True
    set_section_pagination(sec_ch3, fmt="decimal", restart=False)
    # First page footer empty
    sec_ch3.first_page_footer.is_linked_to_previous = False
    sec_ch3.first_page_footer.paragraphs[0].text = ""
    attach_footer_page_number(sec_ch3.footer)
    
    print("Generating Chapter 3: User Interface & Implementation Details...")
    build_chapter_3(doc)
    
    # -------------------------------------------------------------
    # SECTION 6: CHAPTER 4 (Arabic Numerals: Continues from Ch 3)
    # -------------------------------------------------------------
    sec_ch4 = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
    set_page_margins(sec_ch4)
    sec_ch4.different_first_page_header_footer = True
    set_section_pagination(sec_ch4, fmt="decimal", restart=False)
    # First page footer empty
    sec_ch4.first_page_footer.is_linked_to_previous = False
    sec_ch4.first_page_footer.paragraphs[0].text = ""
    attach_footer_page_number(sec_ch4.footer)
    
    print("Generating Chapter 4: Discussion, Evaluation & Future Scope...")
    build_chapter_4(doc)
    
    # -------------------------------------------------------------
    # SECTION 7: REFERENCES & APPENDICES
    # -------------------------------------------------------------
    sec_post = doc.add_section(docx.enum.section.WD_SECTION.NEW_PAGE)
    set_page_margins(sec_post)
    sec_post.different_first_page_header_footer = True
    set_section_pagination(sec_post, fmt="decimal", restart=False)
    sec_post.first_page_footer.is_linked_to_previous = False
    sec_post.first_page_footer.paragraphs[0].text = ""
    attach_footer_page_number(sec_post.footer)
    
    print("Generating References and Appendices...")
    build_references_and_appendices(doc)
    
    print(f"Saving compiled report to {OUTPUT_DOCX}...")
    doc.save(OUTPUT_DOCX)
    print("Report generation complete!")

if __name__ == "__main__":
    main()
