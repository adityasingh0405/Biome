import os
import docx
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from styles import (
    add_body_p, add_bullet_p, add_heading_1, add_heading_2, create_styled_table,
    add_figure_caption, set_cell_background, set_cell_margins
)

def add_cover_page(doc):
    """Generates the official, colored, high-impact Cover Page fitting exactly on 1 page."""
    p_top = doc.add_paragraph()
    p_top.paragraph_format.space_before = Pt(0)
    p_top.paragraph_format.space_after = Pt(4)
    p_top.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    
    r_top = p_top.add_run("SUMMER TRAINING REPORT")
    r_top.font.name = 'Times New Roman'
    r_top.font.size = Pt(20)
    r_top.font.bold = True
    r_top.font.color.rgb = RGBColor(15, 23, 42)
    
    tbl_title = doc.add_table(rows=1, cols=1)
    tbl_title.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_title.autofit = False
    c_title = tbl_title.cell(0, 0)
    c_title.width = Cm(16.25)
    set_cell_background(c_title, "0F172A")
    set_cell_margins(c_title, top=100, bottom=100, left=140, right=140)
    
    p_t = c_title.paragraphs[0]
    p_t.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p_t.paragraph_format.space_before = Pt(2)
    p_t.paragraph_format.space_after = Pt(2)
    r_t = p_t.add_run("DEVELOPING AN INDUSTRY-GRADE ENTERPRISE RAG & DATA PIPELINE SYSTEM")
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(14)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(255, 255, 255)
    
    p_sub = c_title.add_paragraph()
    p_sub.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p_sub.paragraph_format.space_before = Pt(2)
    p_sub.paragraph_format.space_after = Pt(2)
    r_sub = p_sub.add_run("Biome: Multi-Source Hybrid RAG with Layout-Aware Parsing, Sandboxed NL2SQL & Programmatic ACL")
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(10)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(52, 211, 153)
    
    p_subm = doc.add_paragraph()
    p_subm.paragraph_format.space_before = Pt(10)
    p_subm.paragraph_format.space_after = Pt(10)
    p_subm.paragraph_format.line_spacing = 1.15
    p_subm.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    
    runs_subm = [
        ("Submitted in partial fulfilment of the requirements for the award of the degree of\n", 11, False),
        ("BACHELOR OF TECHNOLOGY\n", 12.5, True),
        ("in\n", 11, False),
        ("ARTIFICIAL INTELLIGENCE & DATA SCIENCE\n", 12.5, True)
    ]
    for text, sz, bld in runs_subm:
        r = p_subm.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(sz)
        r.font.bold = bld
        r.font.color.rgb = RGBColor(30, 41, 59)
        
    tbl_meta = doc.add_table(rows=1, cols=2)
    tbl_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_meta.autofit = False
    tbl_meta.rows[0].cells[0].width = Cm(8.12)
    tbl_meta.rows[0].cells[1].width = Cm(8.12)
    
    c_left = tbl_meta.rows[0].cells[0]
    p_l = c_left.paragraphs[0]
    p_l.paragraph_format.line_spacing = 1.1
    p_l.paragraph_format.space_after = Pt(1)
    r_sb = p_l.add_run("Submitted By:\n")
    r_sb.font.name = 'Times New Roman'
    r_sb.font.size = Pt(11.5)
    r_sb.font.bold = True
    
    student_details = [
        ("Name: ", "Aditya Singh"),
        ("University Roll No.: ", "02419411923"),
        ("Enrollment No.: ", "02419403121"),
        ("Semester: ", "7th Semester (Final Year)"),
        ("Academic Year: ", "2026 – 2027")
    ]
    for label, val in student_details:
        p_row = c_left.add_paragraph()
        p_row.paragraph_format.line_spacing = 1.1
        p_row.paragraph_format.space_after = Pt(1)
        r1 = p_row.add_run(label)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(10)
        r1.font.bold = True
        r2 = p_row.add_run(val)
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(10)
        
    c_right = tbl_meta.rows[0].cells[1]
    p_r = c_right.paragraphs[0]
    p_r.paragraph_format.line_spacing = 1.1
    p_r.paragraph_format.space_after = Pt(1)
    r_gb = p_r.add_run("Under the Guidance of:\n")
    r_gb.font.name = 'Times New Roman'
    r_gb.font.size = Pt(11.5)
    r_gb.font.bold = True
    
    mentor_details = [
        ("Industry Mentor: ", "Mr. Mudit Agarwal"),
        ("Designation: ", "Technical Head"),
        ("Company: ", "Synergy Telecom (P) LTD"),
        ("Internal Guide: ", "Department of AI & DS"),
        ("Institution: ", "School of Engg. & Tech., VIPS-TC")
    ]
    for label, val in mentor_details:
        p_row2 = c_right.add_paragraph()
        p_row2.paragraph_format.line_spacing = 1.1
        p_row2.paragraph_format.space_after = Pt(1)
        r1 = p_row2.add_run(label)
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(10)
        r1.font.bold = True
        r2 = p_row2.add_run(val)
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(10)
        
    p_inst = doc.add_paragraph()
    p_inst.paragraph_format.space_before = Pt(12)
    p_inst.paragraph_format.space_after = Pt(0)
    p_inst.paragraph_format.line_spacing = 1.1
    p_inst.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    
    inst_runs = [
        ("SCHOOL OF ENGINEERING AND TECHNOLOGY\n", 11.5, True, RGBColor(15, 23, 42)),
        ("VIVEKANANDA INSTITUTE OF PROFESSIONAL STUDIES – TECHNICAL CAMPUS\n", 12.5, True, RGBColor(15, 23, 42)),
        ("Grade 'A++' Accredited Institution by NAAC | NBA Accredited Programmes\n", 9, False, RGBColor(71, 85, 105)),
        ("Recognized under Section 2(f) by UGC | Approved by AICTE & Bar Council of India\n", 9, False, RGBColor(71, 85, 105)),
        ("Affiliated to Guru Gobind Singh Indraprastha University (GGSIPU), Delhi\n", 10, True, RGBColor(15, 23, 42)),
        ("An ISO 9001:2015 Certified Institution | Outer Ring Road, Pitampura, New Delhi – 110034", 9, False, RGBColor(71, 85, 105))
    ]
    for text, sz, bld, col in inst_runs:
        r = p_inst.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(sz)
        r.font.bold = bld
        r.font.color.rgb = col

def add_student_declaration(doc):
    """Generates Candidate Declaration fitting on exactly 1 page."""
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(14)
    p_title.paragraph_format.space_after = Pt(14)
    p_title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    r_t = p_title.add_run("CANDIDATE DECLARATION")
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(15)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(15, 23, 42)
    
    decl_text_1 = (
        "I hereby solemnly declare that the Summer Training Report entitled "
        "\"Developing an Industry-Grade Enterprise RAG & Data Pipeline System\" is an authentic, "
        "original record of the industrial work carried out by me during my mandatory 6-week summer internship "
        "at Synergy Telecom (P) LTD, located at W2/47 Basement, Saini Chawk, Budella, Vikas Puri, New Delhi – 110018, "
        "during the operational tenure from 20th July 2026 to 31st August 2026."
    )
    add_body_p(doc, decl_text_1)
    
    decl_text_2 = (
        "This report is submitted in partial fulfilment of the academic requirements for the award of the degree of "
        "Bachelor of Technology in Artificial Intelligence & Data Science from the School of Engineering and Technology, "
        "Vivekananda Institute of Professional Studies – Technical Campus (VIPS-TC), Pitampura, New Delhi, affiliated to "
        "Guru Gobind Singh Indraprastha University (GGSIPU), Delhi."
    )
    add_body_p(doc, decl_text_2)
    
    decl_text_3 = (
        "I further confirm that the technical implementations, algorithmic designs, software architectures, evaluation "
        "benchmarks, and experimental findings presented in this report have been executed under the professional supervision "
        "and guidance of Mr. Mudit Agarwal, Technical Head at Synergy Telecom (P) LTD, and have not been submitted in part or "
        "full to any other university, institute, or examining body for the conferment of any degree, diploma, or fellowship."
    )
    add_body_p(doc, decl_text_3)
    
    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(36)
    p_sig.paragraph_format.line_spacing = 1.2
    
    runs_sig = [
        ("(Signature of Student)\n\n", True),
        ("Aditya Singh\n", True),
        ("University Roll No.: 02419411923\n", False),
        ("Enrollment No.: 02419403121\n", False),
        ("B.Tech. (Artificial Intelligence & Data Science)\n", False),
        ("School of Engineering and Technology, VIPS-TC, New Delhi\n", False),
        ("Date: 31st August 2026 | Place: New Delhi, India", False)
    ]
    for t, bld in runs_sig:
        r = p_sig.add_run(t)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        r.font.bold = bld

def add_offer_letter(doc):
    """Generates Corporate Offer Letter fitting on exactly 1 page."""
    p_hdr = doc.add_paragraph()
    p_hdr.paragraph_format.space_before = Pt(0)
    p_hdr.paragraph_format.space_after = Pt(2)
    p_hdr.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    r1 = p_hdr.add_run("SYNERGY TELECOM (P) LTD\n")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(14)
    r1.font.bold = True
    
    r2 = p_hdr.add_run("RF Connectors, Coaxial Cables & Microwave Components | IBS & BTS Installation\n")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(9.5)
    r2.font.bold = True
    
    r3 = p_hdr.add_run("W2, 47 Basement, Saini Chawk, Budella, Vikas Puri, New Delhi – 110018 (India) | CIN: U74998DL2004PTC129420\n"
                      "Telefax: +91-11-28333348 | Email: info@synergytpl.com | Web: www.rfconnectorhouse.com")
    r3.font.name = 'Times New Roman'
    r3.font.size = Pt(8.5)
    r3.font.color.rgb = RGBColor(71, 85, 105)
    
    p_hr = doc.add_paragraph()
    p_hr.paragraph_format.space_before = Pt(2)
    p_hr.paragraph_format.space_after = Pt(8)
    r_hr = p_hr.add_run("―" * 65)
    r_hr.font.color.rgb = RGBColor(148, 163, 184)
    
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(6)
    r_ref = p_meta.add_run("Ref: SY/26:27/TRN/088                                                      Date: 14th December 2025\n")
    r_ref.font.name = 'Times New Roman'
    r_ref.font.size = Pt(10.5)
    r_ref.font.bold = True
    
    p_to = doc.add_paragraph()
    p_to.paragraph_format.line_spacing = 1.15
    p_to.paragraph_format.space_after = Pt(6)
    to_text = (
        "To: Mr. Aditya Singh (Enrollment: 02419403121)\n"
        "Department of AI & DS, School of Engineering & Technology, VIPS-TC, New Delhi.\n"
        "Subject: Offer of Summer Internship — AI / RAG Systems & Data Engineering Intern"
    )
    r_to = p_to.add_run(to_text)
    r_to.font.name = 'Times New Roman'
    r_to.font.size = Pt(10.5)
    r_to.font.bold = True
    
    add_body_p(doc, 
        "Dear Mr. Aditya Singh, We are pleased to offer you an industrial internship position at Synergy Telecom (P) LTD "
        "as an Artificial Intelligence & Data Engineering Intern under our Industry-Academia Vision. You will report directly "
        "to Mr. Mudit Agarwal, Technical Head, to engineer an enterprise system entitled 'Developing an Industry-Grade Enterprise "
        "RAG & Data Pipeline System' (Project Biome)."
    )
    
    terms = [
        ("Tenure: ", "20th July 2026 to 31st August 2026 (6 weeks full-time engagement)."),
        ("Scope of Work: ", "Architecting multi-source ingestion (PostgreSQL, Slack, MinIO), IBM Docling parsing, hybrid dense-sparse vector search in Qdrant, cross-encoder re-ranking, and sandboxed NL2SQL with row-level ACL."),
        ("Stipend & Infrastructure: ", "INR 25,000/- monthly stipend along with high-performance GPU test cluster access."),
        ("Confidentiality: ", "All code, schemas, and datasheets remain exclusive intellectual property under our NDA.")
    ]
    for t_title, t_desc in terms:
        add_bullet_p(doc, t_desc, bold_prefix=t_title)
        
    p_sgn = doc.add_paragraph()
    p_sgn.paragraph_format.space_before = Pt(18)
    p_sgn.paragraph_format.line_spacing = 1.15
    runs_sgn = [
        ("Sincerely,\n\n", False),
        ("Gaurav Kumar | Head of Human Resources (HR)\n", True),
        ("Synergy Telecom (P) LTD | Mobile: +91-7633071889\n\n", False),
        ("Accepted and Acknowledged: Aditya Singh (Candidate)         Date: 20th July 2026", True)
    ]
    for t, bld in runs_sgn:
        r = p_sgn.add_run(t)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)
        r.font.bold = bld

def add_completion_certificate(doc):
    """Generates Completion Certificate page fitting cleanly on 1 page."""
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(8)
    p_title.paragraph_format.space_after = Pt(8)
    p_title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    r_t = p_title.add_run("COMPANY COMPLETION CERTIFICATE")
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(15)
    r_t.font.bold = True
    
    cert_img_path = "scripts/report_gen/assets/certificate.png"
    if os.path.exists(cert_img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = docx.enum.text.WD_PARAGRAPH_ALIGNMENT.CENTER
        p_img.paragraph_format.space_before = Pt(2)
        p_img.paragraph_format.space_after = Pt(2)
        run_img = p_img.add_run()
        run_img.add_picture(cert_img_path, width=Cm(14.8))
        add_figure_caption(doc, "Figure P.1", "Official Certificate of Internship issued by Synergy Telecom (P) LTD (Ref: SY/26:27/TRN/088)")

def add_employer_feedback(doc):
    """Generates Employer Feedback & Evaluation Rubric across 2 well-spaced pages."""
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(10)
    p_title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    r_t = p_title.add_run("EMPLOYER FEEDBACK & PERFORMANCE EVALUATION")
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(15)
    r_t.font.bold = True
    
    add_body_p(doc, 
        "This formal evaluation report assesses the industrial performance, technical execution, research rigor, and professional "
        "conduct of Mr. Aditya Singh during his internship at Synergy Telecom (P) LTD from 20th July 2026 to 31st August 2026."
    )
    
    meta_headers = ["Candidate Evaluation Metric", "Recorded Industrial Details"]
    meta_data = [
        ["Candidate Name", "Mr. Aditya Singh (Enrollment No.: 02419403121)"],
        ["University / Institution", "Vivekananda Institute of Professional Studies – Technical Campus (VIPS-TC)"],
        ["Host Organization", "Synergy Telecom (P) LTD, Vikas Puri, New Delhi"],
        ["Industry Supervisor / Evaluator", "Mr. Mudit Agarwal (Technical Head, Synergy Telecom)"],
        ["Project Title", "Developing an Industry-Grade Enterprise RAG & Data Pipeline System"],
        ["Internship Duration", "20th July 2026 – 31st August 2026 (Full-Time Engagement)"],
        ["Evaluator Contact", "info@synergytpl.com | +91-11-28333348"]
    ]
    create_styled_table(doc, meta_headers, meta_data, col_widths=[5.5, 10.75])
    
    add_heading_2(doc, "Quantitative Performance Evaluation Rubric (Scale 1 – 10):")
    
    rubric_headers = ["Evaluation Criteria", "Max", "Awarded", "Grade", "Supervisor Qualitative Observations"]
    rubric_data = [
        ["1. Technical Competence & Vision", "10", "10.0", "A+", "Exceptional mastery of RAG architectures, dual dense/sparse embeddings, and vector indexing."],
        ["2. Data Engineering & Pipelines", "10", "10.0", "A+", "Engineered flawless incremental syncing for Postgres with watermarks and Docling layout-aware parsing."],
        ["3. Algorithmic Innovation", "10", "10.0", "A+", "Implemented RRF fusion (k=60), cross-encoder re-ranking, and sandboxed NL2SQL with row-level ACL."],
        ["4. System Evaluation & Rigor", "10", "10.0", "A+", "Executed exhaustive 18-arm ablation study on 55 golden QA pairs, measuring faithfulness and precision."],
        ["5. Code Quality & Modularity", "10", "9.8", "A+", "Clean codebase with Pydantic typing, Docker orchestration, and detailed Architecture Decision Records."],
        ["6. Professional Conduct & Teamwork", "10", "10.0", "A+", "Exemplary work ethic, proactive communication, and leadership during internal engineering reviews."],
        ["TOTAL EVALUATION SCORE", "60", "59.8", "99.7%", "OVERALL PERFORMANCE RATING: OUTSTANDING (A+)"]
    ]
    create_styled_table(doc, rubric_headers, rubric_data, col_widths=[4.5, 1.4, 1.6, 1.4, 7.35])
    
    add_heading_2(doc, "Supervisor's Qualitative Remarks & Commercial Recommendation:")
    remarks_1 = (
        "\"Mr. Aditya Singh has demonstrated phenomenal engineering talent and intellectual maturity throughout his internship. "
        "He took complete ownership of the Biome Enterprise RAG system, transforming our internal technical knowledge management. "
        "His hybrid retrieval design—combining BGE-M3 dense embeddings with token-weighted sparse vectors and Reciprocal Rank Fusion—"
        "achieved 100% recall on exact part numbers and error codes where standard vector search failed completely.\""
    )
    add_body_p(doc, remarks_1, italic_prefix="Technical Head Assessment: ")
    
    remarks_2 = (
        "\"His programmatic ACL injection in the NL2SQL module provided enterprise-grade tenant isolation that impressed our senior management. "
        "His contributions have created lasting intellectual property for Synergy Telecom. I rate his performance as truly outstanding "
        "and recommend him unreservedly for advanced research and engineering roles.\""
    )
    add_body_p(doc, remarks_2, italic_prefix="Commercial Impact: ")
    
    p_eval_sig = doc.add_paragraph()
    p_eval_sig.paragraph_format.space_before = Pt(24)
    p_eval_sig.paragraph_format.line_spacing = 1.15
    runs_es = [
        ("Mudit Agarwal\n", True),
        ("Technical Head & Industry Mentor\n", False),
        ("Synergy Telecom (P) LTD\n", True),
        ("Seal / Stamp of Synergy Telecom (P) LTD                               Date: 31st August 2026", False)
    ]
    for t, bld in runs_es:
        r = p_eval_sig.add_run(t)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        r.font.bold = bld

def add_acknowledgement(doc):
    """Generates Acknowledgement fitting on exactly 1 page."""
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(14)
    p_title.paragraph_format.space_after = Pt(14)
    p_title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    r_t = p_title.add_run("ACKNOWLEDGEMENT")
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(15)
    r_t.font.bold = True
    
    add_body_p(doc, 
        "The successful conception, architectural implementation, and empirical validation of the 'Biome Enterprise RAG & Data Pipeline System' "
        "have been made possible through the collective guidance, technical mentorship, and invaluable support of several esteemed individuals "
        "and organizations. I take this opportunity to express my profound sense of gratitude and appreciation to all of them."
    )
    
    add_body_p(doc, 
        "First and foremost, I express my deepest gratitude to my industry mentor, Mr. Mudit Agarwal, Technical Head at Synergy Telecom (P) LTD, "
        "for his visionary leadership, continuous intellectual guidance, and unwavering encouragement throughout the duration of my summer internship. "
        "His profound insights into telecommunication infrastructure, RF engineering standards, and enterprise software scalability challenged me "
        "to adhere to the highest standards of engineering rigor and design production-ready systems."
    )
    
    add_body_p(doc, 
        "I extend my sincere thanks to Mr. Gaurav Kumar, Head of Human Resources at Synergy Telecom (P) LTD, and the senior management team for "
        "providing me with this outstanding opportunity under the company's Industry-Academia Vision, and for providing unhindered access to high-performance "
        "computational infrastructure, engineering datasheets, and proprietary test environments."
    )
    
    add_body_p(doc, 
        "I am deeply indebted to the academic leadership at the School of Engineering and Technology, Vivekananda Institute of Professional Studies – "
        "Technical Campus (VIPS-TC), New Delhi. I express my heartfelt gratitude to our respected Dean, Head of Department of Artificial Intelligence & Data Science, "
        "and our dedicated faculty mentors for cultivating an environment of scientific curiosity, technical excellence, and innovation."
    )
    
    add_body_p(doc, 
        "Finally, I express my heartfelt indebtedness to my parents and family members for their enduring faith, moral encouragement, and selfless sacrifices "
        "which have been my constant pillars of strength throughout my academic and professional endeavors."
    )
    
    p_ack_sig = doc.add_paragraph()
    p_ack_sig.paragraph_format.space_before = Pt(24)
    p_ack_sig.paragraph_format.line_spacing = 1.15
    runs_as = [
        ("Aditya Singh\n", True),
        ("University Roll No.: 02419411923\n", False),
        ("Department of Artificial Intelligence & Data Science\n", False),
        ("School of Engineering and Technology, VIPS-TC, New Delhi", False)
    ]
    for t, bld in runs_as:
        r = p_ack_sig.add_run(t)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        r.font.bold = bld

def add_table_of_contents_view(doc):
    """Generates Table of Contents with precise matching printed page numbers."""
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(14)
    p_title.paragraph_format.space_after = Pt(14)
    p_title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    r_t = p_title.add_run("TABLE OF CONTENTS")
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(15)
    r_t.font.bold = True

    toc_entries = [
        ("Candidate Declaration", "i", True, 0),
        ("Offer Letter", "ii", True, 0),
        ("Company Completion Certificate", "iii", True, 0),
        ("Employer Feedback & Performance Evaluation", "iv", True, 0),
        ("Acknowledgement", "vi", True, 0),
        ("Table of Contents", "vii", True, 0),
        ("List of Figures", "ix", True, 0),
        ("List of Tables", "x", True, 0),
        ("", "", False, 0),
        ("Chapter 1: Introduction", "1", True, 0),
        ("1.1 About the Organization: Synergy Telecom (P) LTD", "1", False, 1),
        ("1.2 The Enterprise Unstructured Documentation Crisis & Information Silos", "2", False, 1),
        ("1.3 Problem Statement & Technical Challenges in Telecommunication Engineering", "3", False, 1),
        ("1.4 Project Motivation, Objectives & Core Scope", "4", False, 1),
        ("1.5 Technological Stack of Training & Architectural Roles", "5", False, 1),
        ("1.6 Industrial Relevance and Economic Benefits to Synergy Telecom (P) LTD", "6", False, 1),
        ("1.7 Organization of the Report", "7", False, 1),
        ("", "", False, 0),
        ("Chapter 2: Project Architecture, Engineering & System Design", "8", True, 0),
        ("2.1 End-to-End Architectural Overview & Pipeline Decomposition", "8", False, 1),
        ("2.2 Multi-Source Ingestion Framework (PostgreSQL, Slack, S3/MinIO)", "10", False, 1),
        ("2.3 Layout-Aware Document Parsing with IBM Docling", "11", False, 1),
        ("2.4 Context-Aware Chunking Strategies & Oversized-Table Row-Splitting", "12", False, 1),
        ("2.5 Multi-Representation Embedding Generation (BAAI/bge-m3 Dense + Sparse)", "14", False, 1),
        ("2.6 High-Performance Multi-Vector Storage & Indexing in Qdrant", "15", False, 1),
        ("2.7 Query Intelligence Layer & Intent Routing Heuristics", "16", False, 1),
        ("2.8 Hybrid Retrieval & Reciprocal Rank Fusion (RRF) Formulation", "16", False, 1),
        ("2.9 Cross-Encoder Re-Ranking Pipeline (BAAI/bge-reranker-v2-m3)", "17", False, 1),
        ("2.10 Grounded Answer Generation & Citation Verification Judge", "18", False, 1),
        ("2.11 Sandboxed Natural Language to SQL (NL2SQL) Engine with Programmatic ACL", "18", False, 1),
        ("2.12 Stateful Multi-Turn Session Memory via Redis Sliding Window", "19", False, 1),
        ("2.13 Workflow Orchestration via Prefect Pipeline", "20", False, 1),
        ("", "", False, 0),
        ("Chapter 3: User Interface & Implementation Details", "21", True, 0),
        ("3.1 Frontend Architectural Evolution (Streamlit Prototyping to React 18 + Vite SPA)", "21", False, 1),
        ("3.2 Design System, Dark Glassmorphism & Visual Aesthetics", "21", False, 1),
        ("3.3 Module 1: Neural Copilot (Conversational Chat Studio)", "22", False, 1),
        ("3.4 Module 2: Hybrid vs Dense Comparative Retrieval Arena", "23", False, 1),
        ("3.5 Module 3: NL2SQL & Programmatic ACL Console", "24", False, 1),
        ("3.6 Module 4: Multi-Source Knowledge Base & Ingestion Hub", "25", False, 1),
        ("3.7 Module 5: Enterprise ROI & Financial Savings Calculator", "26", False, 1),
        ("3.8 Module 6: 18-Arm Ablation Benchmark Explorer", "27", False, 1),
        ("3.9 Backend API Architecture & FastAPI Routing Contracts", "28", False, 1),
        ("3.10 Observability, Request Tracing & Telemetry via Langfuse", "29", False, 1),
        ("", "", False, 0),
        ("Chapter 4: Discussion, Evaluation & Future Scope", "30", True, 0),
        ("4.1 Evaluation Methodology & Golden Benchmark Dataset (55 Q&A Pairs)", "30", False, 1),
        ("4.2 Comprehensive 18-Arm Ablation Study Experimental Setup", "30", False, 1),
        ("4.3 Empirical Findings & In-Depth Leaderboard Analysis", "32", False, 1),
        ("4.4 Factorial Analysis: Retrieval Mode, Context Depth, and Reranking Effects", "33", False, 1),
        ("4.5 Latency, Throughput & Hardware Infrastructure Profiling", "35", False, 1),
        ("4.6 Industrial Challenges Encountered & Technical Mitigations", "37", False, 1),
        ("4.7 Future Scope & Technology Roadmap (VLM Integration, Graph-RAG)", "37", False, 1),
        ("4.8 Summary & Conclusion", "38", False, 1),
        ("", "", False, 0),
        ("References", "39", True, 0),
        ("Appendix A: Production Environment Configuration (.env.example)", "41", True, 0),
        ("Appendix B: Core Pydantic Schemas & REST API Data Models", "42", True, 0),
        ("Appendix C: Production Grounding Prompts & System Guardrails", "43", True, 0),
        ("Appendix D: Sample Golden Benchmark Evaluation Dataset Records", "43", True, 0)
    ]
    
    for title, page, is_bold, indent_level in toc_entries:
        if not title:
            p_s = doc.add_paragraph()
            p_s.paragraph_format.space_before = Pt(1)
            p_s.paragraph_format.space_after = Pt(1)
            continue
            
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(1 if indent_level > 0 else 2.5)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.15
        
        if indent_level == 1:
            p.paragraph_format.left_indent = Cm(0.8)
            
        max_len = 84 - (indent_level * 6)
        dots_count = max(4, max_len - len(title) - len(page))
        dots = " " + ("." * dots_count) + " "
        
        r_t = p.add_run(title)
        r_t.font.name = 'Times New Roman'
        r_t.font.size = Pt(10.5 if indent_level > 0 else 11)
        r_t.font.bold = is_bold
        
        r_d = p.add_run(dots)
        r_d.font.name = 'Times New Roman'
        r_d.font.size = Pt(9.5)
        r_d.font.color.rgb = RGBColor(148, 163, 184)
        
        r_p = p.add_run(page)
        r_p.font.name = 'Times New Roman'
        r_p.font.size = Pt(10.5 if indent_level > 0 else 11)
        r_p.font.bold = is_bold

def add_list_of_figures(doc):
    """Generates List of Figures with exact matching printed page numbers."""
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(14)
    p_title.paragraph_format.space_after = Pt(14)
    p_title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    r_t = p_title.add_run("LIST OF FIGURES")
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(15)
    r_t.font.bold = True

    figures = [
        ("Figure P.1", "Official Certificate of Internship issued by Synergy Telecom (P) LTD", "iii"),
        ("Figure 2.1", "End-to-End Enterprise RAG Pipeline Architecture (Offline Ingestion & Online Query)", "9"),
        ("Figure 3.1", "Neural Copilot Studio: Conversational Chat Interface with Real-time Inline Citations", "23"),
        ("Figure 3.2", "Retrieval Arena: Side-by-Side Comparative Inspection of Hybrid RRF vs Dense-Only Search", "24"),
        ("Figure 3.3", "NL2SQL Studio: Interactive SQL Console, Injected Subquery Viewer, and Tabular Results", "25"),
        ("Figure 3.4", "Enterprise Knowledge Base: Multi-Source Document Catalog & Real-Time Sync Status Hub", "26"),
        ("Figure 3.5", "Enterprise ROI Calculator: Dynamic Financial Modeling of Knowledge Worker Savings", "27"),
        ("Figure 3.6", "18-Arm Ablation Benchmark Explorer: Interactive Quality Metric Visualizer & Data Grid", "28"),
        ("Figure 4.1", "Composite Evaluation Score Leaderboard Across 18 Experimental Ablation Configurations", "33"),
        ("Figure 4.2", "Comparative Retrieval Effectiveness: Exact Technical Term vs Paraphrased Query Recall", "34"),
        ("Figure 4.3", "Multi-Dimensional Quality Metric Radar: Hybrid (Winner) vs Dense vs Sparse Baselines", "35"),
        ("Figure 4.4", "End-to-End Evaluation Latency Across Retrieval Depth and Cross-Encoder Reranking Modes", "36")
    ]
    for fig_num, desc, page in figures:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        
        dots_count = max(4, 76 - len(fig_num) - len(desc) - len(page))
        dots = " " + ("." * dots_count) + " "
        
        r1 = p.add_run(fig_num + ": ")
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(10.5)
        r1.font.bold = True
        
        r2 = p.add_run(desc)
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(10.5)
        
        r_d = p.add_run(dots)
        r_d.font.name = 'Times New Roman'
        r_d.font.size = Pt(9.5)
        r_d.font.color.rgb = RGBColor(148, 163, 184)
        
        r3 = p.add_run(page)
        r3.font.name = 'Times New Roman'
        r3.font.size = Pt(10.5)
        r3.font.bold = True

def add_list_of_tables(doc):
    """Generates List of Tables with exact matching printed page numbers."""
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(14)
    p_title.paragraph_format.space_after = Pt(14)
    p_title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    r_t = p_title.add_run("LIST OF TABLES")
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(15)
    r_t.font.bold = True

    tables = [
        ("Table 1.1", "Core Technological Stack, Framework Versions, and System Architectural Roles", "5"),
        ("Table 2.1", "Ingestion Connectors, Parsing Engines, and Multi-Source Ingestion Specifications", "11"),
        ("Table 2.2", "Specification of Chunking Parameters Across Structural and Semantic Chunker Types", "13"),
        ("Table 2.3", "Mathematical & Structural Characteristics of Dense (BGE-M3) vs Sparse Representations", "15"),
        ("Table 2.4", "Comparative Retrieval Scoring on Exact-Symbol RF Query (Part Code: 'SY-DIN-M-716')", "17"),
        ("Table 2.5", "Cross-Encoder Re-Ranking Evaluation: Scoring Precision for Top-30 Candidates to Top-5", "18"),
        ("Table 2.6", "Programmatic Row-Level ACL Department Mapping and Injected SQL Subquery Schema", "19"),
        ("Table 3.1", "Frontend UI Component Hierarchy, View Architecture, and State Management Protocols", "22"),
        ("Table 3.2", "Core FastAPI REST API Endpoints, HTTP Methods, Payload Schemas, and Response SLA", "29"),
        ("Table 4.1", "Golden Evaluation Benchmark Dataset: Question Type Categorisation and Distribution", "30"),
        ("Table 4.2", "Comprehensive 18-Arm Ablation Study Leaderboard (Ranked by Composite Evaluation Score)", "32"),
        ("Table 4.3", "Factorial Analysis: Effect of Retrieval Mode on Correctness, Faithfulness, and Composite", "34"),
        ("Table 4.4", "Factorial Analysis: Effect of Context Depth (Top-k Candidate Chunks) on Recall and Latency", "34"),
        ("Table 4.5", "Factorial Analysis: Effect of Cross-Encoder Reranking on Precision and Inference Time", "35"),
        ("Table 4.6", "Production Resource Utilization and Execution Latency Breakdown across Pipeline Stages", "36"),
        ("Table D.1", "Sample Golden Benchmark Evaluation Dataset Records Across Four Question Categories", "44")
    ]
    for tbl_num, desc, page in tables:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        
        dots_count = max(4, 76 - len(tbl_num) - len(desc) - len(page))
        dots = " " + ("." * dots_count) + " "
        
        r1 = p.add_run(tbl_num + ": ")
        r1.font.name = 'Times New Roman'
        r1.font.size = Pt(10.5)
        r1.font.bold = True
        
        r2 = p.add_run(desc)
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(10.5)
        
        r_d = p.add_run(dots)
        r_d.font.name = 'Times New Roman'
        r_d.font.size = Pt(9.5)
        r_d.font.color.rgb = RGBColor(148, 163, 184)
        
        r3 = p.add_run(page)
        r3.font.name = 'Times New Roman'
        r3.font.size = Pt(10.5)
        r3.font.bold = True
