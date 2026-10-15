import docx
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

# Academic Margins as specified
MARGIN_LEFT_CM = 3.5
MARGIN_TOP_CM = 2.5
MARGIN_RIGHT_CM = 1.25
MARGIN_BOTTOM_CM = 1.25

def set_page_margins(section):
    """Applies strict A4 dimensions and required margins."""
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(MARGIN_LEFT_CM)
    section.top_margin = Cm(MARGIN_TOP_CM)
    section.right_margin = Cm(MARGIN_RIGHT_CM)
    section.bottom_margin = Cm(MARGIN_BOTTOM_CM)

def configure_document_styles(doc):
    """Configures global document styles according to formatting requirements."""
    styles = doc.styles

    # Normal (Body Text) Style
    normal_style = styles['Normal']
    normal_font = normal_style.font
    normal_font.name = 'Times New Roman'
    normal_font.size = Pt(12)
    normal_font.color.rgb = RGBColor(0, 0, 0)
    
    # Paragraph format for body
    normal_pf = normal_style.paragraph_format
    normal_pf.line_spacing = 1.5
    normal_pf.space_after = Pt(4)
    normal_pf.space_before = Pt(0)
    normal_pf.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY

def add_page_number(run):
    """Inserts a dynamic Word PAGE field into a text run."""
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    run._r.append(fldSimple)

def set_section_pagination(section, fmt="decimal", start_num=None, restart=False):
    """Sets section page number formatting and restart behavior."""
    sectPr = section._sectPr
    pgNumTypes = sectPr.xpath('./w:pgNumType')
    if pgNumTypes:
        pgNumType = pgNumTypes[0]
    else:
        pgNumType = OxmlElement('w:pgNumType')
        sectPr.append(pgNumType)
    
    if fmt:
        pgNumType.set(qn('w:fmt'), fmt)
    
    if restart and start_num is not None:
        pgNumType.set(qn('w:start'), str(start_num))
    elif not restart:
        if qn('w:start') in pgNumType.attrib:
            del pgNumType.attrib[qn('w:start')]

def add_heading_1(doc, text, is_chapter=False):
    """Adds a Heading 1 element in Times New Roman, bold, formatted."""
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.5
    if is_chapter:
        p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    else:
        p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(15 if not is_chapter else 16)
    run.font.bold = True
    run.font.color.rgb = RGBColor(15, 23, 42)
    return p

def add_heading_2(doc, text):
    """Adds a Heading 2 element in Times New Roman, bold, formatted."""
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.5
    p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = RGBColor(30, 41, 59)
    return p

def add_heading_3(doc, text):
    """Adds a Heading 3 element in Times New Roman, bold italic, formatted."""
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.5
    p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    
    run = p.add_run(text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = RGBColor(51, 65, 85)
    return p

def add_body_p(doc, text, bold_prefix=None, italic_prefix=None):
    """Adds a body paragraph with standard 1.5 line spacing, 12pt Times New Roman."""
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    
    if bold_prefix:
        r_bold = p.add_run(bold_prefix)
        r_bold.font.name = 'Times New Roman'
        r_bold.font.size = Pt(12)
        r_bold.font.bold = True
        
    if italic_prefix:
        r_it = p.add_run(italic_prefix)
        r_it.font.name = 'Times New Roman'
        r_it.font.size = Pt(12)
        r_it.font.italic = True

    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    return p

def add_bullet_p(doc, text, bold_prefix=None):
    """Adds a bullet point with hanging indent."""
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    
    if bold_prefix:
        rb = p.add_run(bold_prefix)
        rb.font.name = 'Times New Roman'
        rb.font.size = Pt(12)
        rb.font.bold = True
        
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    return p

def add_figure_caption(doc, fig_num_text, caption_text):
    """Adds a formal figure caption BELOW the figure."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_with_next = False
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    
    r1 = p.add_run(fig_num_text + ": ")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(11)
    r1.font.bold = True
    
    r2 = p.add_run(caption_text)
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(11)
    r2.font.italic = True
    return p

def add_table_caption(doc, tbl_num_text, caption_text):
    """Adds a formal table caption ABOVE the table."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.keep_with_next = True
    p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
    
    r1 = p.add_run(tbl_num_text + ": ")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(11)
    r1.font.bold = True
    
    r2 = p.add_run(caption_text)
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(11)
    r2.font.bold = True
    return p

def set_cell_background(cell, fill_hex):
    """Fills table cell with specified hex background colour."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    """Sets cell internal padding in dxa."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4"):
    """Applies neat, clean academic borders to table."""
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:left w:val="none"/>'
        f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:right w:val="none"/>'
        f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'  <w:insideV w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

def create_styled_table(doc, headers, data, col_widths=None):
    """Constructs a beautifully styled, high-impact academic table."""
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="94A3B8", sz="6")
    
    # Format Header Row
    hdr_row = table.rows[0]
    # Repeat header on every page in Word
    trPr = hdr_row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    
    for i, h_text in enumerate(headers):
        cell = hdr_row.cells[i]
        set_cell_background(cell, "1E293B") # Dark Slate
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        p = cell.paragraphs[0]
        p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(h_text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        
    # Format Data Rows
    for r_idx, row_vals in enumerate(data):
        row = table.rows[r_idx + 1]
        bg_color = "F8FAFC" if (r_idx % 2 == 1) else "FFFFFF"
        for c_idx, val in enumerate(row_vals):
            cell = row.cells[c_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            p = cell.paragraphs[0]
            p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            r = p.add_run(str(val))
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(15, 23, 42)
            
    # Apply column widths if provided
    if col_widths:
        for r in table.rows:
            for c_idx, w in enumerate(col_widths):
                r.cells[c_idx].width = Cm(w)
                
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(8)
    return table

def add_callout_box(doc, title, text_lines):
    """Adds a highlighted callout box for equations or key architectural highlights."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Cm(16.25)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    
    # Left accent border
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'  <w:top w:val="none"/>'
        f'  <w:left w:val="single" w:sz="24" w:space="0" w:color="0D9488"/>'
        f'  <w:bottom w:val="none"/>'
        f'  <w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    r_t = p.add_run(title)
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(11)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(15, 118, 110)
    
    for line in text_lines:
        p_l = cell.add_paragraph()
        p_l.paragraph_format.space_before = Pt(2)
        p_l.paragraph_format.space_after = Pt(2)
        p_l.paragraph_format.line_spacing = 1.15
        r_l = p_l.add_run(line)
        r_l.font.name = 'Times New Roman'
        r_l.font.size = Pt(10.5)
        r_l.font.italic = True
        r_l.font.color.rgb = RGBColor(30, 41, 59)
        
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)
