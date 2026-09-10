import os
from faker import Faker
from docx import Document
from pptx import Presentation

fake = Faker()

# 1. Define the local data bucket directory tree
BUCKET_DIR = "./local_data_bucket"
UNSTRUCTURED_DIR = os.path.join(BUCKET_DIR, "unstructured")

# Create the folders automatically
os.makedirs(UNSTRUCTURED_DIR, exist_ok=True)

def generate_mock_txt(filename, paragraphs_count=3):
    path = os.path.join(UNSTRUCTURED_DIR, filename)
    with open(path, "w", encoding="utf-8") as f:
        f.write(f"# {filename.replace('_', ' ').title()}\n\n")
        for _ in range(paragraphs_count):
            f.write(fake.paragraph(nb_sentences=5) + "\n\n")
    print(f"-> Created TXT file: {filename}")

def generate_mock_docx(filename):
    path = os.path.join(UNSTRUCTURED_DIR, filename)
    doc = Document()
    doc.add_heading(filename.replace('.docx', '').replace('_', ' ').title(), 0)
    
    doc.add_heading('1. Overview & Objectives', level=1)
    doc.add_paragraph(fake.paragraph(nb_sentences=4))
    
    doc.add_heading('2. Standard Operating Procedures', level=1)
    doc.add_paragraph(fake.paragraph(nb_sentences=6))
    
    doc.save(path)
    print(f"-> Created DOCX file: {filename}")

def generate_mock_pptx(filename):
    path = os.path.join(UNSTRUCTURED_DIR, filename)
    prs = Presentation()
    
    # Slide 1: Title Slide
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = "Enterprise Strategy Review"
    slide.placeholders[1].text = f"Prepared by: {fake.company()} - Confidential"
    
    # Slide 2: Bullet points Slide
    bullet_layout = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(bullet_layout)
    slide2.shapes.title.text = "Key Operational Metrics"
    tf = slide2.placeholders[1].text_frame
    tf.text = "Highlights for the current review period:"
    for _ in range(3):
        p = tf.add_paragraph()
        p.text = fake.sentence(nb_words=6)
        p.level = 1
        
    prs.save(path)
    print(f"-> Created PPTX file: {filename}")

if __name__ == "__main__":
    print(f"Initializing local data bucket at: {os.path.abspath(BUCKET_DIR)}")
    print("Populating unstructured files...")
    
    # Generate a rich set of mock files mimicking an enterprise file server
    generate_mock_txt("hr_remote_work_policy.txt", paragraphs_count=4)
    generate_mock_txt("it_security_guidelines.txt", paragraphs_count=5)
    generate_mock_docx("company_onboarding_manual.docx")
    generate_mock_docx("quarterly_compliance_report.docx")
    generate_mock_pptx("q3_logistics_overview.pptx")
    
    print("\nLocal data bucket is ready and filled with files!")
    print(f"Files are located at: {os.path.abspath(UNSTRUCTURED_DIR)}")